# -*- coding: utf-8 -*-
"""Gera a minuta do contrato de consultoria R2 x Surcin e o termo de confidencialidade (30/10/2026).

Uso: python gerar_contrato.py <pasta de saída>
Campos entre colchetes ficam para preencher. É minuta: precisa de revisão jurídica antes de assinar.
"""
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

SAIDA = sys.argv[1] if len(sys.argv) > 1 else "."
AZUL = RGBColor(0x1F, 0x3A, 0x5F)
CINZA = RGBColor(0x59, 0x59, 0x59)


def novo_doc(titulo, subtitulo):
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2.5)
        s.top_margin = s.bottom_margin = Cm(2.2)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    p = d.add_paragraph()
    r = p.add_run(titulo)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = AZUL
    p = d.add_paragraph()
    r = p.add_run(subtitulo)
    r.italic = True
    r.font.color.rgb = CINZA
    aviso(d, "MINUTA PARA REVISÃO. Os campos entre colchetes precisam ser preenchidos. Recomenda-se revisão "
             "jurídica antes da assinatura. As condições seguem a Proposta de consultoria, versão 7, de 30/10/2026.")
    return d


def aviso(d, texto):
    p = d.add_paragraph()
    r = p.add_run(texto)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x9C, 0x5B, 0x00)


def clausula(d, titulo):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run(titulo)
    r.bold = True
    r.font.color.rgb = AZUL


def item(d, texto):
    p = d.add_paragraph(texto)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def assinaturas(d, cidade_data="[Cidade], [data]"):
    d.add_paragraph()
    d.add_paragraph(cidade_data)
    for nome in ["CONTRATANTE: [razão social da Surcin] · Thamires Surcin",
                 "CONTRATADA: RENATO ROMAO DA SILVA LTDA (R2 Tech Consultoria) · Renato Romão da Silva",
                 "Testemunha 1: [nome e CPF]",
                 "Testemunha 2: [nome e CPF]"]:
        d.add_paragraph("\n______________________________________________")
        d.add_paragraph(nome)


def contrato():
    d = novo_doc("Contrato de prestação de serviços de consultoria",
                 "R2 Tech Consultoria e Surcin Assessoria Empresarial · minuta de 30/10/2026")

    clausula(d, "Partes")
    item(d, "CONTRATANTE: [razão social], inscrita no CNPJ sob o nº [__], com sede em [endereço], neste ato "
            "representada por Thamires Surcin, [qualificação]. [Confirmar qual dos CNPJs da Surcin contrata.]")
    item(d, "CONTRATADA: RENATO ROMAO DA SILVA LTDA, que atua sob o nome R2 Tech Consultoria, sociedade "
            "empresária limitada inscrita no CNPJ sob o nº 50.563.506/0001-87, com sede na Rua Siqueira Campos, 253, "
            "apto 901, Copacabana, Rio de Janeiro/RJ, CEP 22031-071, neste ato representada por seu sócio-administrador, "
            "Renato Romão da Silva, [nacionalidade, estado civil, CPF].")

    clausula(d, "Cláusula 1 · Objeto")
    item(d, "1.1. Prestação de serviços de consultoria em organização, processos, rotinas do Claude e pilotos de "
            "automação assistida, conforme a Proposta de consultoria, versão 7, de 30/10/2026 (Anexo I), que "
            "integra este contrato.")
    item(d, "1.2. O serviço é contratado por inteiro e executado em cinco etapas: organização (Etapas 1 e 2, cerca "
            "de 70 horas) e implantação (Etapas 3 a 5, cerca de 151 horas).")
    item(d, "1.3. O escopo é fechado nos limites da proposta: até 5 entrevistas, 8 processos mapeados, 5 fluxos "
            "redesenhados, 5 rotinas do Claude, 2 pilotos, 2 integrações, 3 treinamentos de até 90 minutos, "
            "1 painel por área e da diretoria e 1 ciclo de avaliação por competência.")
    item(d, "1.4. Não fazem parte do objeto os itens listados como fora do escopo na proposta, em especial "
            "decisões e protocolos profissionais, pagamentos e disparos automáticos, auditoria fiscal ou contábil "
            "formal, revisão trabalhista e suporte contínuo.")

    clausula(d, "Cláusula 2 · Prazo e cronograma")
    item(d, "2.1. As Etapas 1 e 2 vão de 03/11/2026 a 04/12/2026. As Etapas 3 a 5 vão de 04/01/2027 a "
            "12/02/2027, data da entrega final.")
    item(d, "2.2. Cada etapa termina com um aceite por escrito da CONTRATANTE. Sem manifestação em [5] dias úteis "
            "após a entrega, a etapa é considerada aceita. [Confirmar o prazo.]")
    item(d, "2.3. Atrasos causados por falta de acesso, dados, agenda ou retorno da CONTRATANTE deslocam o "
            "cronograma na mesma medida, sem custo adicional para nenhuma das partes.")

    clausula(d, "Cláusula 3 · Preço e pagamento")
    item(d, "3.1. Pelo objeto completo, a CONTRATANTE paga R$ 31.900,00 (trinta e um mil e novecentos reais), "
            "pelo serviço completo.")
    item(d, "3.2. O pagamento é feito em 10 (dez) parcelas mensais de R$ 3.190,00, a primeira em 03/11/2026 e as "
            "demais no dia [03] de cada mês, até agosto de 2027, por [PIX / boleto / transferência] contra nota fiscal.")
    item(d, "3.3. A entrega final ocorre antes do fim do pagamento; as parcelas restantes continuam devidas nas "
            "datas previstas.")
    item(d, "3.4. Em caso de atraso, incidem multa de [2]% e juros de [1]% ao mês sobre a parcela em atraso. "
            "[Confirmar percentuais.]")
    item(d, "3.5. Licenças, planos de ferramentas, uso de modelos de IA, integrações, hospedagem e serviços de "
            "terceiros são pagos diretamente pela CONTRATANTE, mediante aprovação prévia.")

    clausula(d, "Cláusula 4 · Obrigações da CONTRATADA")
    item(d, "4.1. Executar os serviços com diligência, nas datas combinadas, e entregar os materiais em formato "
            "editável, no ambiente da CONTRATANTE.")
    item(d, "4.2. Avisar com antecedência riscos de atraso ou de mudança de escopo.")
    item(d, "4.3. Usar acessos apenas para o objeto deste contrato, no menor nível necessário, e devolvê-los ou "
            "pedir sua revogação ao fim do contrato.")
    item(d, "4.4. Cumprir o Termo de confidencialidade e proteção de dados (Anexo II).")

    clausula(d, "Cláusula 5 · Obrigações da CONTRATANTE")
    item(d, "5.1. Indicar um ponto focal para acessos e agenda, liberar os acessos de leitura necessários e "
            "disponibilizar a equipe para as entrevistas.")
    item(d, "5.2. Responder e validar as entregas nos prazos combinados.")
    item(d, "5.3. Fornecer, para os pilotos, dados autorizados e, sempre que possível, anonimizados.")
    item(d, "5.4. Manter as decisões profissionais, fiscais, trabalhistas e jurídicas, que continuam sob "
            "responsabilidade da CONTRATANTE e de seus profissionais habilitados.")

    clausula(d, "Cláusula 6 · Uso de inteligência artificial")
    item(d, "6.1. As rotinas do Claude preparam, comparam e sinalizam. Uma pessoa da CONTRATANTE confere a "
            "evidência, decide e assume a responsabilidade. Nenhuma rotina paga, transmite obrigações, envia "
            "mensagens a clientes ou decide matéria fiscal, trabalhista ou jurídica sem confirmação humana.")
    item(d, "6.2. A CONTRATADA não responde por decisões tomadas sem a conferência prevista no item 6.1.")

    clausula(d, "Cláusula 7 · Propriedade dos materiais")
    item(d, "7.1. Os materiais produzidos para a CONTRATANTE passam a ser dela após a quitação de cada parcela.")
    item(d, "7.2. A CONTRATADA mantém a titularidade de sua metodologia pré-existente (Arquitetura da Ordem™), "
            "modelos genéricos e ferramentas próprias, e concede à CONTRATANTE licença de uso interno, por prazo "
            "indeterminado, do que for incorporado aos materiais entregues.")
    item(d, "7.3. A CONTRATADA só pode citar a CONTRATANTE como cliente ou caso com autorização por escrito.")

    clausula(d, "Cláusula 8 · Mudanças de escopo")
    item(d, "8.1. Qualquer item além do escopo fechado é orçado à parte e só executado após aditivo ou aprovação "
            "por escrito.")

    clausula(d, "Cláusula 9 · Rescisão")
    item(d, "9.1. Qualquer parte pode rescindir este contrato com aviso por escrito de [15] dias. São devidos os "
            "valores das etapas já entregues ou em andamento, proporcionalmente às horas executadas. "
            "[Confirmar regra.]")
    item(d, "9.2. Na rescisão, a CONTRATADA entrega os materiais produzidos até a data e encerra os acessos.")

    clausula(d, "Cláusula 10 · Foro")
    item(d, "10.1. Fica eleito o foro da comarca de [__] para dirimir questões deste contrato.")

    item(d, "Anexo I · Proposta de consultoria, versão 7, de 30/10/2026.")
    item(d, "Anexo II · Termo de confidencialidade e proteção de dados.")
    assinaturas(d)
    d.save(f"{SAIDA}/contrato-consultoria-surcin-minuta.docx")


def termo():
    d = novo_doc("Termo de confidencialidade e proteção de dados",
                 "Anexo II do contrato R2 Tech Consultoria e Surcin · minuta de 30/10/2026")

    clausula(d, "1 · Informação confidencial")
    item(d, "É confidencial toda informação da CONTRATANTE e de seus clientes a que a CONTRATADA tiver acesso: "
            "dados cadastrais, fiscais, financeiros, trabalhistas e jurídicos, documentos, senhas, processos "
            "internos, informações sobre a equipe e sobre a situação da empresa.")

    clausula(d, "2 · Obrigações")
    item(d, "2.1. Usar a informação apenas para executar o contrato.")
    item(d, "2.2. Não copiar, divulgar ou compartilhar a informação com terceiros sem autorização por escrito.")
    item(d, "2.3. Acessar os sistemas por convite em conta própria, no menor nível necessário, sem usar senhas "
            "compartilhadas ou listas de senhas da CONTRATANTE.")
    item(d, "2.4. Não retirar dados de clientes do ambiente da CONTRATANTE. Para testes e pilotos, usar dados "
            "anonimizados ou previamente autorizados.")
    item(d, "2.5. Informar a CONTRATANTE em até [24 horas] sobre qualquer incidente de segurança que envolva "
            "seus dados.")

    clausula(d, "3 · Proteção de dados pessoais (Lei nº 13.709/2018)")
    item(d, "3.1. No tratamento de dados pessoais de clientes e colaboradores da CONTRATANTE, a CONTRATANTE atua "
            "como controladora e a CONTRATADA como operadora, seguindo as instruções da CONTRATANTE.")
    item(d, "3.2. A CONTRATADA adota medidas de segurança adequadas e só envia dados a ferramentas de IA ou de "
            "terceiros aprovadas pela CONTRATANTE.")
    item(d, "3.3. Ao fim do contrato, a CONTRATADA devolve ou elimina os dados recebidos, salvo obrigação legal "
            "de guarda.")

    clausula(d, "4 · Vigência")
    item(d, "As obrigações deste termo valem durante o contrato e por [5] anos após o seu término. "
            "[Confirmar prazo.]")

    clausula(d, "5 · Descumprimento")
    item(d, "O descumprimento sujeita a parte responsável a reparar os danos comprovados, sem prejuízo das "
            "demais medidas legais.")
    assinaturas(d)
    d.save(f"{SAIDA}/termo-confidencialidade-surcin-minuta.docx")


if __name__ == "__main__":
    contrato()
    termo()
    print("ok")
