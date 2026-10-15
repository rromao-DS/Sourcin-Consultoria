# -*- coding: utf-8 -*-
"""Gera o deck da apresentação da proposta Surcin (versão preliminar)."""
import sys
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SAIDA = sys.argv[1] if len(sys.argv) > 1 else "."
AZUL = RGBColor(0x1F, 0x3A, 0x5F)
CLARO = RGBColor(0xE8, 0xEE, 0xF5)
CINZA = RGBColor(0x59, 0x59, 0x59)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
AMARELO = RGBColor(0xFF, 0xF4, 0xD6)

prs = Presentation()
prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
VAZIO = prs.slide_layouts[6]


def caixa(slide, x, y, w, h, texto="", tam=16, cor=CINZA, negrito=False, fundo=None):
    if fundo is not None:
        s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
        s.fill.solid()
        s.fill.fore_color.rgb = fundo
        s.line.fill.background()
        s.adjustments[0] = 0.08
        s.shadow.inherit = False
    else:
        s = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = Cm(0.5)
    linhas = texto if isinstance(texto, list) else [texto]
    for i, l in enumerate(linhas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = l
        r.font.size = Pt(tam)
        r.font.color.rgb = cor
        r.font.bold = negrito or (fundo is not None and i == 0)
        p.alignment = PP_ALIGN.LEFT
        r.font.name = "Calibri"
        p.space_after = Pt(6)
    return s


def slide_titulo(titulo, subtitulo=None):
    s = prs.slides.add_slide(VAZIO)
    faixa = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Cm(0.5))
    faixa.fill.solid()
    faixa.fill.fore_color.rgb = AZUL
    faixa.line.fill.background()
    caixa(s, 1.8, 1.3, 30, 2, titulo, 30, AZUL, True)
    if subtitulo:
        caixa(s, 1.8, 3.2, 30, 1.2, subtitulo, 16, CINZA)
    caixa(s, 1.8, 17.6, 30, 0.8, "Surcin · Proposta de consultoria · Confidencial", 9, CINZA)
    return s


# 1 · Capa
s = prs.slides.add_slide(VAZIO)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = AZUL
bg.line.fill.background()
caixa(s, 2.5, 5.5, 29, 2.5, "Organização, estrutura e processos", 40, BRANCO, True)
caixa(s, 2.5, 8.3, 29, 1.5, "Proposta de consultoria para a Surcin Assessoria Empresarial", 22, BRANCO)
caixa(s, 2.5, 13.5, 29, 1.2, "Renato Romão · R2 Tech Consultoria · 15 de outubro de 2026", 16, BRANCO)

# 2 · O que eu entendi
s = slide_titulo("O que eu entendi", "\"Eu tenho um processo. Ele só não está eficiente.\"")
caixa(s, 1.8, 4.8, 9.6, 11.5, ["A Surcin hoje",
                               "• 40 a 45 empresas no fiscal",
                               "• Cerca de 12 a 16 clientes no financeiro",
                               "• DP com ~25 vidas",
                               "• ~20 processos jurídicos legados",
                               "• Júlia, Vitória, Tamara e Alana",
                               "• Trello, OneDrive, Veri, Domínio, Conta Azul"], 14, AZUL, fundo=CLARO)
caixa(s, 12.1, 4.8, 9.6, 11.5, ["O que dói",
                                "• Tudo passa por você",
                                "• Fiscal nas suas mãos desde agosto",
                                "• Parâmetros fiscais errados: ISS, multa, cliente perdido",
                                "• Fechamento financeiro concentrado no fim do mês",
                                "• Guia na pasta errada, tarefa concluída sem evidência"], 14, AZUL, fundo=AMARELO)
caixa(s, 22.4, 4.8, 9.6, 11.5, ["O que você quer",
                                "• Equipe com autonomia",
                                "• Escalar com as mesmas pessoas",
                                "• Tempo para a prova de 13/12",
                                "• Entrar em 2027 pronta para contratar e vender"], 14, BRANCO, fundo=AZUL)

# 3 · Como vamos trabalhar
s = slide_titulo("Como vamos trabalhar", "Processo e controle antes de ferramenta. Cada fase termina com uma entrega e uma decisão sua.")
etapas = [("Fase 0", "Reunião e briefing", "Concluída, sem custo"),
          ("Fase 1", "Contenção, mapeamento, estrutura e plano", "Nesta proposta"),
          ("Fase 2", "Implantação em módulos", "Orçada após o plano"),
          ("Fase 3", "Acompanhamento mensal", "Opcional")]
for i, (n, nome, st) in enumerate(etapas):
    x = 1.8 + i * 7.6
    cor = AZUL if st == "Nesta proposta" else CLARO
    txt = BRANCO if st == "Nesta proposta" else AZUL
    caixa(s, x, 6, 7.2, 7, [n, nome, "", st], 15, txt, fundo=cor)

# 3b · Antes da prova
s = slide_titulo("Antes da sua prova", "Três controles manuais em uso até 15/11, sem software e sem depender de você no dia a dia")
ctrl = [("Guia conferida", "Cliente, CNPJ, pasta, tributo, competência e valor conferidos antes de pagar"),
        ("Evidência para concluir", "Cartão crítico do Trello só fecha com o comprovante no OneDrive"),
        ("Dúvida registrada", "Toda dúvida levada a você vira registro: mede a dependência e alimenta o futuro assistente")]
for i, (t, d) in enumerate(ctrl):
    caixa(s, 1.8 + i * 10.2, 5.5, 9.6, 8, [t, "", d], 16, AZUL, fundo=CLARO)

# 4 · Estrutura organizacional
s = slide_titulo("Estrutura organizacional", "Hipótese para validarmos juntas: organizar pelo caminho do cliente")
caixa(s, 12.4, 4.5, 9, 1.8, ["Direção · Thamires"], 16, BRANCO, True, fundo=AZUL)
areas = [("Revisão técnica", "Contador(a) em 2027 [decisão dela]"), ("Entrada e relacionamento", "Júlia"),
         ("Financeiro", "Vitória"), ("Departamento Pessoal", "Tamara"), ("Jurídico", "Legado: a decidir")]
for i, (a, p) in enumerate(areas):
    caixa(s, 1.8 + i * 6.1, 7.6, 5.7, 2.6, [a, p], 13, AZUL, fundo=CLARO)
caixa(s, 1.8, 11.3, 30, 5, ["Entregável: organograma atual e proposto, papéis, responsabilidades e indicadores por função.",
                            "É a base para delegar revisões, receber o(a) contador(a) e criar a remuneração por pontos."],
      15, CINZA)

# 5 · Processos
s = slide_titulo("Processos identificados", "11 processos; a prioridade inicial é onde há mais risco e mais tempo gasto")
caixa(s, 1.8, 4.8, 14.5, 11.5, ["Prioridade 1",
                                "• Fiscal mensal (44 clientes, risco de multa)",
                                "• Conciliação e fechamento financeiro",
                                "", "Prioridade 2",
                                "• Implantação de novos clientes",
                                "• Atendimento e triagem",
                                "• Emissão de notas"], 15, BRANCO, fundo=AZUL)
caixa(s, 17.5, 4.8, 14.5, 11.5, ["Prioridade 3",
                                 "• Pagamentos dos clientes do financeiro",
                                 "• Folha e DP",
                                 "• Comercial",
                                 "• Abertura de empresa",
                                 "• Acompanhamento jurídico",
                                 "• Reinf, IRPF e parcelamentos",
                                 "", "Escolha final na primeira semana, com você."], 15, AZUL, fundo=CLARO)

# 6 · Duas opções
s = slide_titulo("Duas opções de escopo", "Você escolhe o tamanho; a Essencial pode virar Completa depois")
caixa(s, 1.8, 5, 14.5, 10.8, ["Essencial · 19/10 a ~22/11",
                              "• Entrevistas com 4 pessoas",
                              "• Até 5 processos mapeados",
                              "• Gargalos e dependências",
                              "• Organograma e matriz de responsabilidades",
                              "• Lista priorizada de automações"], 15, AZUL, fundo=CLARO)
caixa(s, 17.5, 5, 14.5, 10.8, ["Completa (recomendada) · 19/10 a 22/01",
                               "• 3 controles de contenção até 15/11",
                               "• Entrevistas com até 5 pessoas",
                               "• Até 8 processos mapeados",
                               "• Gargalos e dependências",
                               "• Organograma, papéis e indicadores",
                               "• 5 fluxos otimizados",
                               "• Plano de automação com esforço e custo",
                               "• Avaliação de desempenho por competência (jan.)"], 15, BRANCO, fundo=AZUL)

# 7 · Cronograma
s = slide_titulo("Cronograma", "Cerca de 9 horas suas no total, sem reuniões na semana dos fechamentos nem na semana da prova")
marcos = [("19/10", "Kick-off e entrevistas"), ("26/10", "Entrevistas com Tamara e Vitória"),
          ("15/11", "Entrega do mapeamento"), ("29/11", "Estrutura organizacional"),
          ("13/12", "Sua prova (sem reuniões)"), ("18/12", "Fluxos e plano de automação"),
          ("22/01", "Avaliação por competência e entrega final")]
for i, (dt, t) in enumerate(marcos):
    fundo = AMARELO if dt == "13/12" else (AZUL if dt in ("15/11", "22/01") else CLARO)
    cor = BRANCO if fundo == AZUL else AZUL
    caixa(s, 1.8 + i * 4.33, 6, 4.1, 5.5, [dt, t], 14, cor, fundo=fundo)

# 8 · Investimento
s = slide_titulo("Investimento", "Preço fechado. A implementação é orçada depois, por pacotes.")
precos = [("Essencial", "R$ 7.800", "50% no aceite, 50% na entrega"),
          ("Completa (recomendada)", "R$ 17.400", "30% no aceite, 35% no mapeamento, 35% na entrega final")]
for i, (t, v, d) in enumerate(precos):
    fundo = AZUL if i == 1 else CLARO
    cor = BRANCO if i == 1 else AZUL
    caixa(s, 1.8 + i * 15.6, 5.3, 14.6, 6.5, [t, v, d], 20, cor, True, fundo=fundo)
caixa(s, 1.8, 12.6, 30, 4, ["A reunião inicial e esta proposta não foram cobradas.",
                            "A Completa inclui a avaliação de desempenho por competência, feita em janeiro de 2027.",
                            "Fora da Fase 1: desenvolvimento e integrações (Fase 2), correção do Domínio, migração de contas, valores da premiação e licenças."],
      14, CINZA)

# 8b · Fase 2
s = slide_titulo("Depois: implantação em módulos", "Cada módulo é orçado após o plano, com prova de viabilidade e pagamento sempre humano")
mods = ["Ambiente Surcin e Base Mestre", "Especialista de dúvidas", "Validador de guias",
        "Fechamento interativo (piloto)", "Auditoria Trello × OneDrive", "Fila de pedidos de pagamento",
        "Painel de desempenho", "Comercial e CRM (2027)"]
for i, m in enumerate(mods):
    caixa(s, 1.8 + (i % 4) * 7.6, 5 + (i // 4) * 4.6, 7.2, 3.8, [m], 15, AZUL, fundo=CLARO)
caixa(s, 1.8, 14.6, 30, 2, ["Porte de referência: P R$ 1,5 a 2,8 mil · M R$ 3,5 a 6,5 mil · G R$ 8 a 13 mil. Licenças por conta da Surcin."], 13, CINZA)

# 9 · Próximos passos
s = slide_titulo("Próximos passos")
caixa(s, 1.8, 4.5, 30, 11, ["1. Escolha da opção e aceite da proposta",
                            "2. Acordo de confidencialidade",
                            "3. Kick-off na semana de 19/10, com a Júlia como ponto focal",
                            "", "Para confirmar hoje:",
                            "• Quantos clientes o financeiro atende hoje?",
                            "• Quem emite as notas a partir de novembro?",
                            "• A equipe pode conversar antes das 10h30 ou no almoço?",
                            "• Tráfego volta em 15/12 ou em janeiro?",
                            "• Quem corrige a parametrização do Domínio?"], 18, AZUL)

prs.save(f"{SAIDA}/06-apresentacao-proposta.pptx")
print("ok")
