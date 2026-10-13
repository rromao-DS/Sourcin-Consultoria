# -*- coding: utf-8 -*-
"""Gera a planilha de aplicação da avaliação por competências da Surcin."""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
import modelo as m  # noqa: E402

SAIDA = sys.argv[1] if len(sys.argv) > 1 else "."

AZUL = "1F3A5F"
F_TIT = Font(bold=True, size=16, color=AZUL)
F_SUB = Font(italic=True, size=9, color="595959")
F_CAB = Font(bold=True, color="FFFFFF")
F_NEG = Font(bold=True)
P_CAB = PatternFill("solid", fgColor=AZUL)
P_ENT = PatternFill("solid", fgColor="FFF4D6")   # célula de preenchimento
P_CALC = PatternFill("solid", fgColor="E8EEF5")  # célula calculada
P_SEC = PatternFill("solid", fgColor="D9E1EC")
FINO = Side(style="thin", color="BFBFBF")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
QUEBRA = Alignment(wrap_text=True, vertical="top")
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)

PESSOAS = list(m.FUNCOES.keys())
ABA = {p: p for p in PESSOAS}


def cab(ws, linha, textos, col=1):
    for i, t in enumerate(textos):
        c = ws.cell(row=linha, column=col + i, value=t)
        c.font, c.fill, c.alignment, c.border = F_CAB, P_CAB, CENTRO, BORDA


def cel(ws, ref, valor=None, fill=None, fonte=None, alin=QUEBRA, fmt=None):
    c = ws[ref]
    if valor is not None:
        c.value = valor
    c.border = BORDA
    c.alignment = alin
    if fill:
        c.fill = fill
    if fonte:
        c.font = fonte
    if fmt:
        c.number_format = fmt
    return c


def titulo(ws, texto, sub=m.VERSAO):
    ws["A1"] = texto
    ws["A1"].font = F_TIT
    ws["A2"] = sub
    ws["A2"].font = F_SUB


def larguras(ws, ws_larg):
    for col, w in ws_larg.items():
        ws.column_dimensions[col].width = w


def main():
    wb = Workbook()

    # ------------------------------------------------------------------ Instruções
    ws = wb.active
    ws.title = "Instruções"
    titulo(ws, "Avaliação por competências · Surcin")
    passos = [
        "Proposta para a Thamires validar. Pesos, faixas e valores ficam na aba Parâmetros e podem ser alterados.",
        "Células amarelas são de preenchimento; células azuis são calculadas. Não apague as fórmulas.",
        "Use um arquivo por ciclo (semestre). Antes de começar, preencha o nome do ciclo na aba Parâmetros.",
        "1. Todo mês, lance faltas injustificadas, atrasos e faltas justificadas/atestados na aba Assiduidade.",
        "2. No fim do semestre, cada funcionária preenche a coluna Autoavaliação da sua aba, com exemplos.",
        "3. A Thamires preenche a coluna Avaliação Thamires e os dois indicadores da função (em %).",
        "4. Na conversa individual, fechem a Nota final. Se ficar em branco, vale a nota da Thamires.",
        "5. Registre de 2 a 3 ações na aba PDI (plano de desenvolvimento individual).",
        "6. No 2º semestre, copie a nota final do 1º semestre para a aba Resumo. A nota do ano é a média dos dois.",
        "7. A gratificação só é paga no fim do ano. O valor é limitado à reserva acumulada (aba Parâmetros).",
        "Escala: 1 Não atende · 2 Atende em parte · 3 Atende · 4 Supera · 5 Referência. "
        "A descrição de cada nível está na aba Competências.",
        "Atestados e faltas justificadas são só registrados: não reduzem a nota (ver documento, seção 7).",
        "Confidencial: cada funcionária vê só a própria aba. Guarde o arquivo em pasta de acesso restrito. "
        "Não anexe atestados aqui; registre só a quantidade.",
    ]
    for i, t in enumerate(passos, start=4):
        ws.cell(row=i, column=1, value=t).alignment = QUEBRA
    ws.column_dimensions["A"].width = 120

    # ------------------------------------------------------------------ Parâmetros
    pa = wb.create_sheet("Parâmetros")
    titulo(pa, "Parâmetros")
    cel(pa, "A4", "Ciclo", fonte=F_NEG)
    cel(pa, "B4", "1º semestre de 2027", fill=P_ENT)

    cab(pa, 6, ["Bloco", "Peso"])
    for i, (nome, peso) in enumerate(m.PESOS):
        cel(pa, f"A{7 + i}", nome)
        cel(pa, f"B{7 + i}", peso, fill=P_ENT, fmt="0%")
    cel(pa, "A10", "Soma (deve ser 100%)", fonte=F_NEG)
    cel(pa, "B10", "=SUM(B7:B9)", fill=P_CALC, fmt="0%")

    # tabelas de conversão (limite inferior -> nota)
    cab(pa, 12, ["Faltas injustificadas (a partir de)", "Nota"])
    for i, (lim, nota) in enumerate(m.ASSIDUIDADE[0][1]):
        cel(pa, f"A{13 + i}", lim, fill=P_ENT)
        cel(pa, f"B{13 + i}", nota, fill=P_ENT)
    cab(pa, 12, ["Atrasos (a partir de)", "Nota"], col=4)
    for i, (lim, nota) in enumerate(m.ASSIDUIDADE[1][1]):
        cel(pa, f"D{13 + i}", lim, fill=P_ENT)
        cel(pa, f"E{13 + i}", nota, fill=P_ENT)
    cab(pa, 12, ["Cumprimento do indicador (a partir de)", "Nota"], col=7)
    for i, (lim, nota) in enumerate(m.FAIXAS_PCT):
        cel(pa, f"G{13 + i}", lim, fill=P_ENT, fmt="0%")
        cel(pa, f"H{13 + i}", nota, fill=P_ENT)

    cab(pa, 19, ["Nota do ano (a partir de)", "% da referência", "Faixa"])
    for i, (lim, pct, txt) in enumerate(m.FAIXAS_GRAT):
        cel(pa, f"A{20 + i}", lim, fill=P_ENT, fmt="0.00")
        cel(pa, f"B{20 + i}", pct, fill=P_ENT, fmt="0%")
        cel(pa, f"C{20 + i}", txt.split(":")[0], fill=P_ENT)

    cab(pa, 25, ["Funcionária", "Valor de referência da gratificação (R$)", "Participa da gratificação?"])
    for i, p in enumerate(PESSOAS):
        cel(pa, f"A{26 + i}", p)
        cel(pa, f"B{26 + i}", 0, fill=P_ENT, fmt='"R$" #,##0.00')
        cel(pa, f"C{26 + i}", "A definir" if p == "Alana" else "Sim", fill=P_ENT)
    cel(pa, "A31", "Reserva mensal para o fundo (R$)", fonte=F_NEG)
    cel(pa, "B31", 0, fill=P_ENT, fmt='"R$" #,##0.00')
    cel(pa, "A32", "Meses de reserva no ano", fonte=F_NEG)
    cel(pa, "B32", 12, fill=P_ENT)
    cel(pa, "A33", "Fundo disponível no fim do ano (R$)", fonte=F_NEG)
    cel(pa, "B33", "=B31*B32", fill=P_CALC, fmt='"R$" #,##0.00')
    pa["A35"] = "Valores zerados de propósito: a Thamires define a referência e a reserva."
    pa["A35"].font = F_SUB
    dv_sim = DataValidation(type="list", formula1='"Sim,Não,A definir"', allow_blank=True)
    pa.add_data_validation(dv_sim)
    dv_sim.add("C26:C29")
    larguras(pa, {"A": 42, "B": 26, "C": 26, "D": 22, "E": 8, "G": 26, "H": 8})

    # ------------------------------------------------------------------ Competências (catálogo)
    co = wb.create_sheet("Competências")
    titulo(co, "Catálogo de competências")
    cab(co, 4, ["Código", "Competência", "Definição", "Nível 1", "Nível 3", "Nível 5", "Origem"])
    li = 5
    for cod, nome, defin, n1, n3, n5, orig in m.COMUNS:
        for j, v in enumerate([cod, nome, defin, n1, n3, n5, orig]):
            cel(co, f"{'ABCDEFG'[j]}{li}", v)
        li += 1
    li += 1
    cab(co, li, ["Código", "Funcionária · competência técnica", "O que é atender (nota 3)", "Evidência"])
    li += 1
    for p, f in m.FUNCOES.items():
        for cod, nome, atende, evid in f["tecnicas"]:
            for j, v in enumerate([cod, f"{p} · {nome}", atende, evid]):
                cel(co, f"{'ABCD'[j]}{li}", v)
            li += 1
    li += 1
    cab(co, li, ["Código", "Trilha de supervisão (Júlia)", "O que se espera"])
    li += 1
    for cod, nome, desc in m.LIDERANCA:
        for j, v in enumerate([cod, nome, desc]):
            cel(co, f"{'ABC'[j]}{li}", v)
        li += 1
    larguras(co, {"A": 8, "B": 30, "C": 45, "D": 35, "E": 35, "F": 35, "G": 30})

    # ------------------------------------------------------------------ Assiduidade
    asd = wb.create_sheet("Assiduidade")
    titulo(asd, "Registro mensal de assiduidade")
    asd["A3"] = "=\"Ciclo: \"&Parâmetros!B4"
    cab(asd, 4, ["Mês", "Funcionária", "Faltas injustificadas", "Atrasos", "Faltas justificadas / atestados (só registro)", "Observação"])
    li = 5
    for mes in range(1, 7):
        for p in PESSOAS:
            cel(asd, f"A{li}", f"Mês {mes}", fill=P_ENT)
            cel(asd, f"B{li}", p)
            for col in "CDE":
                cel(asd, f"{col}{li}", 0, fill=P_ENT, alin=CENTRO)
            cel(asd, f"F{li}", None, fill=P_ENT)
            li += 1
    ult_asd = li - 1
    dv_int = DataValidation(type="whole", operator="between", formula1="0", formula2="200", allow_blank=True)
    asd.add_data_validation(dv_int)
    dv_int.add(f"C5:E{ult_asd}")
    larguras(asd, {"A": 12, "B": 14, "C": 14, "D": 10, "E": 22, "F": 40})
    asd.freeze_panes = "A5"

    # ------------------------------------------------------------------ Abas individuais
    dv_nota = None
    resumo_refs = {}
    for idx, (p, f) in enumerate(m.FUNCOES.items()):
        ws = wb.create_sheet(ABA[p])
        titulo(ws, f"{p} · {f['cargo']}")
        ws["A3"] = "=\"Ciclo: \"&Parâmetros!B4"
        ws["A3"].font = F_NEG
        cab(ws, 5, ["Bloco", "Código", "Competência", "Autoavaliação", "Avaliação Thamires",
                    "Nota final (acordada)", "Nota considerada", "Exemplos e evidências do período"])
        dv_nota = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                                 error="Use uma nota de 1 a 5.", errorTitle="Nota inválida")
        ws.add_data_validation(dv_nota)

        def linha_comp(li, bloco, cod, nome):
            cel(ws, f"A{li}", bloco)
            cel(ws, f"B{li}", cod, alin=CENTRO)
            cel(ws, f"C{li}", nome)
            for col in "DEF":
                cel(ws, f"{col}{li}", None, fill=P_ENT, alin=CENTRO)
            dv_nota.add(f"D{li}:F{li}")
            cel(ws, f"G{li}", f'=IF(F{li}<>"",F{li},IF(E{li}<>"",E{li},""))', fill=P_CALC, alin=CENTRO)
            cel(ws, f"H{li}", None, fill=P_ENT)

        li = 6
        ini_c = li
        for cod, nome, *_ in m.COMUNS:
            linha_comp(li, "Comportamental", cod, nome)
            li += 1
        fim_c = li - 1
        ini_t = li
        for cod, nome, *_ in f["tecnicas"]:
            linha_comp(li, "Técnica", cod, nome)
            li += 1
        fim_t = li - 1

        li += 1
        cab(ws, li, ["Bloco", "", "Indicador objetivo", "Valor do período", "", "", "Nota", "Fonte / observação"])
        li += 1
        ini_i = li
        # faltas e atrasos vêm da aba Assiduidade
        cel(ws, f"A{li}", "Indicador")
        cel(ws, f"C{li}", "Faltas injustificadas no semestre")
        cel(ws, f"D{li}", f'=SUMIFS(Assiduidade!$C$5:$C${ult_asd},Assiduidade!$B$5:$B${ult_asd},"{p}")',
            fill=P_CALC, alin=CENTRO)
        cel(ws, f"G{li}", f"=LOOKUP(D{li},Parâmetros!$A$13:$A$17,Parâmetros!$B$13:$B$17)", fill=P_CALC, alin=CENTRO)
        cel(ws, f"H{li}", "Aba Assiduidade")
        li += 1
        cel(ws, f"A{li}", "Indicador")
        cel(ws, f"C{li}", "Atrasos no semestre")
        cel(ws, f"D{li}", f'=SUMIFS(Assiduidade!$D$5:$D${ult_asd},Assiduidade!$B$5:$B${ult_asd},"{p}")',
            fill=P_CALC, alin=CENTRO)
        cel(ws, f"G{li}", f"=LOOKUP(D{li},Parâmetros!$D$13:$D$17,Parâmetros!$E$13:$E$17)", fill=P_CALC, alin=CENTRO)
        cel(ws, f"H{li}", "Aba Assiduidade")
        li += 1
        for nome_ind, fonte in f["indicadores"]:
            cel(ws, f"A{li}", "Indicador")
            cel(ws, f"C{li}", nome_ind)
            cel(ws, f"D{li}", None, fill=P_ENT, alin=CENTRO, fmt="0%")
            cel(ws, f"G{li}", f'=IF(D{li}="","",LOOKUP(D{li},Parâmetros!$G$13:$G$17,Parâmetros!$H$13:$H$17))',
                fill=P_CALC, alin=CENTRO)
            cel(ws, f"H{li}", fonte, fill=P_ENT)
            li += 1
        fim_i = li - 1
        cel(ws, f"A{li}", "Registro")
        cel(ws, f"C{li}", "Faltas justificadas / atestados (não entram na nota)")
        cel(ws, f"D{li}", f'=SUMIFS(Assiduidade!$E$5:$E${ult_asd},Assiduidade!$B$5:$B${ult_asd},"{p}")',
            fill=P_CALC, alin=CENTRO)
        li += 2

        cab(ws, li, ["Resultado", "", "Bloco", "Média", "Peso", "", "", ""])
        li += 1
        r_c, r_t, r_i = li, li + 1, li + 2
        for k, (nome_b, ini, fim, prow) in enumerate([
                (m.PESOS[0][0], ini_c, fim_c, 7), (m.PESOS[1][0], ini_t, fim_t, 8), (m.PESOS[2][0], ini_i, fim_i, 9)]):
            r = li + k
            cel(ws, f"C{r}", nome_b)
            cel(ws, f"D{r}", f'=IFERROR(ROUND(AVERAGE(G{ini}:G{fim}),2),"")', fill=P_CALC, alin=CENTRO, fmt="0.00")
            cel(ws, f"E{r}", f"=Parâmetros!B{prow}", fill=P_CALC, alin=CENTRO, fmt="0%")
        li += 3
        r_f = li
        cel(ws, f"C{r_f}", "Nota final do ciclo", fonte=F_NEG)
        cel(ws, f"D{r_f}",
            f'=IF(OR(D{r_c}="",D{r_t}="",D{r_i}=""),"",ROUND(D{r_c}*E{r_c}+D{r_t}*E{r_t}+D{r_i}*E{r_i},2))',
            fill=P_CALC, alin=CENTRO, fonte=F_NEG, fmt="0.00")
        cel(ws, f"E{r_f}", f'=IF(D{r_f}="","",LOOKUP(D{r_f},Parâmetros!$A$20:$A$23,Parâmetros!$C$20:$C$23))',
            fill=P_CALC)
        ws.merge_cells(f"E{r_f}:H{r_f}")
        resumo_refs[p] = (r_c, r_t, r_i, r_f)

        if p == "Júlia":
            li = r_f + 2
            ws.cell(row=li, column=1, value="Trilha de supervisão: desenvolvimento, não entra na nota nem na gratificação").font = F_NEG
            li += 1
            cab(ws, li, ["Bloco", "Código", "Competência", "Autoavaliação", "Avaliação Thamires", "", "", "Exemplos e próximos passos"])
            li += 1
            for cod, nome, _ in m.LIDERANCA:
                cel(ws, f"A{li}", "Supervisão")
                cel(ws, f"B{li}", cod, alin=CENTRO)
                cel(ws, f"C{li}", nome)
                for col in "DE":
                    cel(ws, f"{col}{li}", None, fill=P_ENT, alin=CENTRO)
                dv_nota.add(f"D{li}:E{li}")
                cel(ws, f"H{li}", None, fill=P_ENT)
                li += 1

        if f.get("obs"):
            ws.cell(row=ws.max_row + 2, column=1, value=f["obs"]).font = F_SUB

        for faixa in (f"G{ini_c}:G{fim_t}", f"G{ini_i}:G{fim_i}"):
            ws.conditional_formatting.add(faixa, CellIsRule(operator="lessThan", formula=["3"],
                                                            fill=PatternFill("solid", fgColor="F8D7DA")))
        larguras(ws, {"A": 15, "B": 8, "C": 44, "D": 14, "E": 16, "F": 14, "G": 12, "H": 50})
        ws.freeze_panes = "A6"

    # ------------------------------------------------------------------ Resumo
    rs = wb.create_sheet("Resumo", 1)
    titulo(rs, "Resumo do ciclo e gratificação")
    rs["A3"] = "=\"Ciclo: \"&Parâmetros!B4"
    cab(rs, 5, ["Funcionária", "Comportamentais", "Técnicas", "Indicadores", "Nota deste ciclo",
                "Nota do 1º semestre (copiar no 2º ciclo)", "Nota do ano", "Faixa", "% da referência",
                "Valor de referência", "Participa?", "Gratificação calculada", "Gratificação ajustada ao fundo"])
    for i, p in enumerate(PESSOAS):
        r = 6 + i
        r_c, r_t, r_i, r_f = resumo_refs[p]
        aba = f"'{ABA[p]}'"
        cel(rs, f"A{r}", p, fonte=F_NEG)
        cel(rs, f"B{r}", f"={aba}!D{r_c}", fill=P_CALC, alin=CENTRO, fmt="0.00")
        cel(rs, f"C{r}", f"={aba}!D{r_t}", fill=P_CALC, alin=CENTRO, fmt="0.00")
        cel(rs, f"D{r}", f"={aba}!D{r_i}", fill=P_CALC, alin=CENTRO, fmt="0.00")
        cel(rs, f"E{r}", f"={aba}!D{r_f}", fill=P_CALC, alin=CENTRO, fmt="0.00")
        cel(rs, f"F{r}", None, fill=P_ENT, alin=CENTRO, fmt="0.00")
        cel(rs, f"G{r}", f'=IF(E{r}="","",IF(F{r}="",E{r},ROUND(AVERAGE(E{r},F{r}),2)))',
            fill=P_CALC, alin=CENTRO, fmt="0.00")
        cel(rs, f"H{r}", f'=IF(G{r}="","",LOOKUP(G{r},Parâmetros!$A$20:$A$23,Parâmetros!$C$20:$C$23))', fill=P_CALC)
        cel(rs, f"I{r}", f'=IF(G{r}="","",LOOKUP(G{r},Parâmetros!$A$20:$A$23,Parâmetros!$B$20:$B$23))',
            fill=P_CALC, alin=CENTRO, fmt="0%")
        cel(rs, f"J{r}", f"=Parâmetros!B{26 + i}", fill=P_CALC, fmt='"R$" #,##0.00')
        cel(rs, f"K{r}", f"=Parâmetros!C{26 + i}", fill=P_CALC, alin=CENTRO)
        cel(rs, f"L{r}", f'=IF(OR(I{r}="",K{r}<>"Sim"),0,ROUND(I{r}*J{r},2))', fill=P_CALC, fmt='"R$" #,##0.00')
        cel(rs, f"M{r}", f"=IF(L$10=0,0,ROUND(L{r}*MIN(1,Parâmetros!$B$33/L$10),2))", fill=P_CALC,
            fmt='"R$" #,##0.00')
    cel(rs, "A10", "Total", fonte=F_NEG)
    cel(rs, "L10", "=SUM(L6:L9)", fill=P_CALC, fonte=F_NEG, fmt='"R$" #,##0.00')
    cel(rs, "M10", "=SUM(M6:M9)", fill=P_CALC, fonte=F_NEG, fmt='"R$" #,##0.00')
    cel(rs, "A11", "Fundo disponível", fonte=F_NEG)
    cel(rs, "M11", "=Parâmetros!B33", fill=P_CALC, fmt='"R$" #,##0.00')
    rs["A13"] = ("A gratificação é apurada só no fim do ano. Se a soma calculada passar do fundo, todas são "
                 "reduzidas na mesma proporção. Valores e forma de pagamento dependem de revisão trabalhista.")
    rs["A13"].font = F_SUB
    larguras(rs, {"A": 14, "B": 15, "C": 11, "D": 12, "E": 12, "F": 18, "G": 11, "H": 18, "I": 11,
                  "J": 14, "K": 11, "L": 15, "M": 17})
    rs.row_dimensions[5].height = 45

    # ------------------------------------------------------------------ PDI
    pdi = wb.create_sheet("PDI")
    titulo(pdi, "Plano de desenvolvimento individual")
    pdi["A3"] = "=\"Ciclo: \"&Parâmetros!B4"
    cab(pdi, 5, ["Funcionária", "Competência a desenvolver", "Ação combinada", "Apoio da empresa", "Prazo",
                 "Situação", "Comentário no acompanhamento"])
    li = 6
    dv_sit = DataValidation(type="list", formula1='"Não iniciada,Em andamento,Concluída"', allow_blank=True)
    pdi.add_data_validation(dv_sit)
    for p in PESSOAS:
        for _ in range(3):
            cel(pdi, f"A{li}", p)
            for col in "BCDEFG":
                cel(pdi, f"{col}{li}", None, fill=P_ENT)
            pdi[f"E{li}"].number_format = "dd/mm/yyyy"
            li += 1
    dv_sit.add(f"F6:F{li - 1}")
    larguras(pdi, {"A": 13, "B": 30, "C": 45, "D": 30, "E": 12, "F": 15, "G": 40})

    os.makedirs(SAIDA, exist_ok=True)
    wb.save(os.path.join(SAIDA, "avaliacao-competencias-surcin.xlsx"))


if __name__ == "__main__":
    main()
    print("ok")
